from dotenv import load_dotenv
load_dotenv()  # reads .env file automatically

import os
import json
import re
import io
import time
import urllib.request
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor, as_completed
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, jsonify
from flask_sqlalchemy import SQLAlchemy
from flask_bcrypt import Bcrypt
from werkzeug.utils import secure_filename
import cloudinary
import cloudinary.uploader
import cloudinary.api
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from datetime import datetime
import pytz
from pptx.oxml.xmlchemy import OxmlElement

# --- Dynamic model selection -------------------------------------------------

AI_REQUEST_TIMEOUT_SECONDS = 20

_GEMINI_SYSTEM_INSTRUCTION = (
    "You are a caring, professional college mentor writing an individualized student progress comment. "
    "Follow the supplied parent-facing reference style: begin exactly with 'Dear Parent,'. In the next sentence, "
    "introduce the student's progress by name and registration number. Describe supported positive academic or "
    "activity highlights before attendance. Only mention subjects below 80%, with exact names and percentages, "
    "using constructive and encouraging wording. Acknowledge stronger subjects only if relevant, and address event "
    "participation and online courses only when those details are recorded, and include the supplied branch-relevant "
    "B.Tech skill recommendation verbatim when present. Write 300 to 800 words across readable paragraphs, "
    "aiming for about 450 words when the supplied record supports it. Use the extra space for practical, relevant "
    "guidance and encouragement rather than invented student history. "
    "Write in the mentor's own warm, professional voice; never identify text as 'faculty guidance' or "
    "'mentor observation'. Keep the tone positive, including any below-80% attendance guidance. "
    "Use only supplied facts; never invent achievements, grades, attendance, risks, family circumstances, or actions. "
    "Do not infer attendance, event/course participation, or gender when data is missing. Do not copy private remarks verbatim. "
    "Write readable paragraphs without headings or bullets and return only a JSON object with "
    '"reg_no" and "comment" fields.'
)

def _get_model_choice():
    """Return the configured provider, model name, and API key."""
    gemini_key = (
        (os.getenv("GEMINI_API_KEY") or "").strip()
        or (os.getenv("GOOGLE_API_KEY") or "").strip()
    )
    openai_key = (os.getenv("OPENAI_API_KEY") or "").strip()
    groq_key = (os.getenv("GROQ_API_KEY") or "").strip()
    choice = (os.getenv("MODEL_SELECTION") or "").strip().upper()

    if not choice:
        choice = (
            "GROQ" if groq_key else
            "GEMINI" if gemini_key else
            "OPENAI" if openai_key else
            "GEMINI"
        )

    if choice == "OPENAI":
        return ("openai", os.getenv("OPENAI_MODEL", "gpt-4o-mini"), openai_key or None)
    if choice == "GROQ":
        return (
            "groq",
            os.getenv("GROQ_MODEL", "openai/gpt-oss-120b"),
            groq_key or None,
        )
    if choice == "GEMINI":
        return (
            "gemini",
            os.getenv("GEMINI_MODEL", "gemini-2.5-flash"),
            gemini_key or None,
        )

    return (choice.lower(), "", None)

PROVIDER, MODEL_NAME, API_KEY = _get_model_choice()
AI_PROVIDER_AVAILABLE = False
AI_PROVIDER_ERROR = None
_gemini_model = None
_chat_client = None

if PROVIDER == "gemini":
    if API_KEY:
        try:
            import google.generativeai as genai
            genai.configure(api_key=API_KEY)
            _gemini_model = genai.GenerativeModel(
                model_name=MODEL_NAME,
                system_instruction=_GEMINI_SYSTEM_INSTRUCTION,
            )
            AI_PROVIDER_AVAILABLE = True
        except Exception as exc:
            print(f"[AI Report] Gemini not available: {exc}")
            AI_PROVIDER_AVAILABLE = False
            AI_PROVIDER_ERROR = f"Gemini client initialization failed ({type(exc).__name__})."
            _gemini_model = None
    else:
        print("[AI Report] Gemini not configured. Set GEMINI_API_KEY or GOOGLE_API_KEY to enable AI comment generation.")
        AI_PROVIDER_ERROR = "Gemini API key is not configured."
elif PROVIDER in {"openai", "groq"}:
    if API_KEY:
        try:
            import openai
            client_options = {
                "api_key": API_KEY,
                "timeout": AI_REQUEST_TIMEOUT_SECONDS,
                "max_retries": 0,
            }
            if PROVIDER == "groq":
                client_options["base_url"] = "https://api.groq.com/openai/v1"
            _chat_client = openai.OpenAI(**client_options)
            AI_PROVIDER_AVAILABLE = True
        except Exception as exc:
            print(f"[AI Report] {PROVIDER.title()} not available: {exc}")
            AI_PROVIDER_AVAILABLE = False
            AI_PROVIDER_ERROR = f"{PROVIDER.title()} client initialization failed ({type(exc).__name__})."
            _chat_client = None
    else:
        api_key_name = "GROQ_API_KEY" if PROVIDER == "groq" else "OPENAI_API_KEY"
        print(f"[AI Report] {PROVIDER.title()} not configured. Set {api_key_name} to enable AI comment generation.")
        AI_PROVIDER_ERROR = f"{PROVIDER.title()} API key is not configured."
else:
    print(f"[AI Report] Unsupported provider: {PROVIDER}")
    AI_PROVIDER_ERROR = f"Unsupported AI provider: {PROVIDER}."

# --- Document generation libraries ------------------------------------------------
try:
    from docx import Document
    from docx.shared import Pt as DocxPt, Inches as DocxInches, RGBColor as DocxRGB
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement as DocxOxmlElement
    DOCX_AVAILABLE = True
except Exception as _de:
    print(f"[AI Report] python-docx not available: {_de}")
    DOCX_AVAILABLE = False

def add_highlight(run, color_hex):
    rPr = run._r.get_or_add_rPr()
    highlight = OxmlElement('a:highlight')
    srgbClr = OxmlElement('a:srgbClr')
    srgbClr.set('val', color_hex)
    highlight.append(srgbClr)
    rPr.append(highlight)

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY') or os.urandom(32)
uri = os.environ.get("DATABASE_URL", "sqlite:///database.db")
if uri.startswith("postgres://"):
    uri = uri.replace("postgres://", "postgresql://", 1)
app.config['SQLALCHEMY_DATABASE_URI'] = uri

# Use QueuePool for better concurrency. NullPool opens a new connection per query,
# which causes connection exhaustion / 500 errors when many students use the app at once.
if 'postgresql' in uri or 'postgres' in uri:
    app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {
        'pool_size': 10,
        'max_overflow': 20,
        'pool_timeout': 30,
        'pool_recycle': 280,
        'pool_pre_ping': True,
        'connect_args': {
            'sslmode': 'require',
            'connect_timeout': 10,
        },
    }
else:
    app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {
        'pool_pre_ping': True,
        'pool_recycle': 300,
    }
app.config['UPLOAD_FOLDER'] = 'static/uploads'
app.config['ALLOWED_EXTENSIONS'] = {'png', 'jpg', 'jpeg', 'gif'}
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

db = SQLAlchemy(app)
bcrypt = Bcrypt(app)

# Models
class User(db.Model):
    username = db.Column(db.String(50), primary_key=True)
    password_hash = db.Column(db.String(100), nullable=False)
    role = db.Column(db.String(20), nullable=False) # 'faculty' or 'student'

class StudentDetail(db.Model):
    reg_num = db.Column(db.String(50), primary_key=True)
    name = db.Column(db.String(100))
    course = db.Column(db.String(100))
    studying_year = db.Column(db.String(20))
    branch = db.Column(db.String(100))
    # JSON strings for dynamic data
    slot_info = db.Column(db.Text, default='[]') # List of slot names
    attendance_data = db.Column(db.Text, default='{}') # {"Slot A": 81}
    marks_data = db.Column(db.Text, default='{}') # {"Slot A": {"model": 20}}
    
    registered_new_course = db.Column(db.String(200))
    online_course = db.Column(db.String(200))
    event_participation = db.Column(db.Text)
    additional_description = db.Column(db.Text)
    results_data = db.Column(db.Text, default='[]')  # [{"subject": "Maths", "grade": "A"}]
    custom_advisory = db.Column(db.Text)
    photo_path = db.Column(db.String(200))
    last_updated = db.Column(db.DateTime, nullable=True)

    def get_attendance(self):
        try: return json.loads(self.attendance_data)
        except: return {}
    
    def get_marks(self):
        try: return json.loads(self.marks_data)
        except: return {}
    
    def get_slots(self):
        try: return json.loads(self.slot_info)
        except: return []

class GlobalAdvisory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    content = db.Column(db.Text, nullable=False, default="All students are advised to pay their 2nd-year tuition fees on time through the Viana Portal.\nAdditionally, kindly upload your recent passport-size photograph to your Viana profile at the earliest, if you have not already done so....")

class GlobalMentorObservation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    content = db.Column(db.Text, nullable=False, default="I personally advised the student to concentrate more on study and skill development. The student is currently attending an online course to improve their technical skills, which is really appreciable.")

class GlobalSettings(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    show_grades_in_ppt = db.Column(db.Boolean, nullable=False, default=False)

# Helper functions
def delete_photo(photo_path):
    if not photo_path:
        return
    if 'cloudinary.com' in photo_path:
        try:
            public_id = "simats_profiles/" + photo_path.split('/')[-1].split('.')[0]
            cloudinary.uploader.destroy(public_id)
        except Exception as e:
            print(f'Failed to delete from Cloudinary: {e}')
    elif os.path.exists(photo_path):
        try:
            os.unlink(photo_path)
        except Exception as e:
            print(f'Failed to delete local photo: {e}')

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def seed_db():
    # Rename the existing seeded faculty login without changing its password.
    previous_faculty = User.query.get('Ashwin@123')
    renamed_faculty = User.query.get('Gnanavel@123')
    if previous_faculty and not renamed_faculty:
        previous_faculty.username = 'Gnanavel@123'
    elif previous_faculty and renamed_faculty:
        print("[Account migration] Both faculty usernames exist; keeping both accounts unchanged.")
    elif not renamed_faculty:
        db.session.add(User(
            username='Gnanavel@123',
            password_hash=bcrypt.generate_password_hash(
                os.getenv('GNANAVEL_INITIAL_PASSWORD', 'welcome')
            ).decode('utf-8'),
            role='faculty'
        ))

    # Faculty: admin@123
    if not User.query.get('admin@123'):
        db.session.add(User(
            username='admin@123',
            password_hash=bcrypt.generate_password_hash(
                os.getenv('ADMIN_INITIAL_PASSWORD', 'welcome')
            ).decode('utf-8'),
            role='faculty'
        ))

    # Remove only the two known demo students previously inserted by seed_db.
    demo_student = StudentDetail.query.filter_by(
        reg_num='24REG01',
        name='Ibrahim',
        course='CSA0708 - Computer Networks',
    ).first()
    if demo_student:
        db.session.delete(demo_student)
        demo_account = User.query.filter_by(username='24REG01', role='student').first()
        if demo_account:
            db.session.delete(demo_account)

    demo_account = User.query.filter_by(username='Ibrahim@123', role='student').first()
    if demo_account and not StudentDetail.query.get('Ibrahim@123'):
        db.session.delete(demo_account)

    db.session.commit()

with app.app_context():
    db.create_all()  # Only creates tables if they don't exist — never wipes data
    try:
        from sqlalchemy import inspect
        inspector = inspect(db.engine)
        columns = [col['name'] for col in inspector.get_columns('student_detail')]

        def add_column_if_missing(column_name, ddl):
            if column_name not in columns:
                db.session.execute(db.text(f'ALTER TABLE student_detail ADD COLUMN {column_name} {ddl}'))
                db.session.commit()
                columns.append(column_name)

        add_column_if_missing('last_updated', 'DATETIME' if db.engine.name != 'postgresql' else 'TIMESTAMP')
        add_column_if_missing('results_data', "TEXT DEFAULT '[]'")
        add_column_if_missing('custom_advisory', 'TEXT')
        add_column_if_missing('mentor_username', 'VARCHAR(50)')
        add_column_if_missing('studying_year', 'VARCHAR(20)')
        add_column_if_missing('branch', 'VARCHAR(100)')
    except Exception as e:
        db.session.rollback()
        print(f"Migration error: {e}")
    seed_db()

# Routes
@app.route('/')
def index():
    if 'user' in session:
        if session['role'] == 'faculty':
            return redirect(url_for('faculty_dashboard'))
        return redirect(url_for('student_dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        user = User.query.filter_by(username=username).first()
        
        if user and bcrypt.check_password_hash(user.password_hash, password):
            session['user'] = user.username
            session['role'] = user.role
            if user.role == 'faculty':
                return redirect(url_for('faculty_dashboard'))
            return redirect(url_for('student_dashboard'))
        flash('Invalid username or password', 'danger')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/change_password', methods=['POST'])
def change_password():
    if 'user' not in session:
        return redirect(url_for('login'))
    
    current_password = request.form['current_password']
    new_password = request.form['new_password']
    
    user = User.query.get(session['user'])
    if user and bcrypt.check_password_hash(user.password_hash, current_password):
        user.password_hash = bcrypt.generate_password_hash(new_password).decode('utf-8')
        db.session.commit()
        flash('Password changed successfully!', 'success')
    else:
        flash('Incorrect current password', 'danger')
    
    return redirect(url_for('faculty_dashboard' if session['role'] == 'faculty' else 'student_dashboard'))

@app.route('/student', methods=['GET', 'POST'])
def student_dashboard():
    if 'user' not in session or session['role'] != 'student':
        return redirect(url_for('login'))
    
    student = StudentDetail.query.get(session['user'])
    if not student:
        student = StudentDetail(reg_num=session['user'], name=session['user'])
        db.session.add(student)
        db.session.commit()

    if request.method == 'POST':
        try:
            student.name = (request.form.get('name') or '').strip() or student.name
            student.course = (request.form.get('course') or '').strip() or student.course

            # Dynamic Slots handling
            slots = [s.strip() for s in request.form.getlist('slot_names[]') if s and s.strip()]
            if not slots:
                slots = student.get_slots() or []
            student.slot_info = json.dumps(slots)

            att_data = {}
            marks_data = {}
            current_attendance = student.get_attendance()
            current_marks = student.get_marks()
            for slot in slots:
                att_val = request.form.get(f'att_{slot}', current_attendance.get(slot, 0))
                att_data[slot] = str(att_val).strip() if att_val is not None else 0
                marks_data[slot] = {
                    'model': (request.form.get(f'model_{slot}') or current_marks.get(slot, {}).get('model', '-')).strip(),
                    'test1': (request.form.get(f'test1_{slot}') or current_marks.get(slot, {}).get('test1', '-')).strip(),
                    'test2': (request.form.get(f'test2_{slot}') or current_marks.get(slot, {}).get('test2', '-')).strip(),
                    'avg': (request.form.get(f'avg_{slot}') or current_marks.get(slot, {}).get('avg', '-')).strip(),
                    'total_marks': (request.form.get(f'total_marks_{slot}') or current_marks.get(slot, {}).get('total_marks', '')).strip(),
                    'course': (request.form.get(f'course_{slot}') or current_marks.get(slot, {}).get('course', '')).strip()
                }

            student.attendance_data = json.dumps(att_data)
            student.marks_data = json.dumps(marks_data)

            if 'studying_year' in request.form:
                student.studying_year = (request.form.get('studying_year') or '').strip()
            if 'branch' in request.form:
                student.branch = (request.form.get('branch') or '').strip()
            student.registered_new_course = (request.form.get('registered_new_course') or '').strip()
            student.online_course = (request.form.get('online_course') or '').strip()
            student.event_participation = (request.form.get('event_participation') or '').strip()
            student.additional_description = (request.form.get('description') or '').strip()

            # Results data — save subject/grade entries from My Results tab
            result_subjects = request.form.getlist('result_subject[]')
            result_grades = request.form.getlist('result_grade[]')
            if result_subjects:  # Only update if results fields were submitted
                results = [{'subject': s.strip(), 'grade': g.strip()} for s, g in zip(result_subjects, result_grades) if s and s.strip()]
                student.results_data = json.dumps(results)

            file = request.files.get('photo')
            if file and file.filename and allowed_file(file.filename):
                if os.environ.get('CLOUDINARY_URL'):
                    try:
                        upload_result = cloudinary.uploader.upload(file, folder="simats_profiles")
                        student.photo_path = upload_result.get('secure_url')
                    except Exception as e:
                        flash(f'Failed to upload photo to cloud storage: {str(e)}', 'warning')
                else:
                    try:
                        filename = secure_filename(f"{session['user']}_{file.filename}")
                        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                        file.save(filepath)
                        student.photo_path = filepath
                    except Exception as e:
                        print(f"Local photo upload warning: {e}")

            # Safe Kolkata timestamp calculation (pure Python, immune to OS pytz missing tzdata)
            from datetime import timedelta
            student.last_updated = datetime.utcnow() + timedelta(hours=5, minutes=30)
            
            db.session.commit()
            flash('Details updated successfully!', 'success')
        except Exception as e:
            db.session.rollback()
            print(f"Student submit error: {e}")
            flash('An error occurred while saving your details. Please try again.', 'danger')

        return redirect(url_for('student_dashboard'))
    
    # Passing current data as dicts
    try:
        results = json.loads(student.results_data) if student.results_data else []
    except:
        results = []
    return render_template('student.html', student=student,
                           attendance=student.get_attendance(),
                           marks=student.get_marks(),
                           slots=student.get_slots(),
                           results=results)

@app.route('/add_student', methods=['POST'])
def add_student():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    reg_num = request.form['reg_num']
    name = request.form['name']
    password = request.form['password']
    
    if User.query.get(reg_num):
        flash('Student with this Registration Number already exists.', 'danger')
    else:
        new_user = User(
            username=reg_num,
            password_hash=bcrypt.generate_password_hash(password).decode('utf-8'),
            role='student'
        )
        new_detail = StudentDetail(reg_num=reg_num, name=name, last_updated=datetime.now(pytz.timezone('Asia/Kolkata')).replace(tzinfo=None))
        db.session.add(new_user)
        db.session.add(new_detail)
        db.session.commit()
        flash(f'Student {name} added successfully!', 'success')
    
    return redirect(url_for('faculty_dashboard'))

@app.route('/edit_student_faculty/<reg_num>', methods=['POST'])
def edit_student_faculty(reg_num):
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    student = StudentDetail.query.get(reg_num)
    if student:
        student.name = request.form.get('name', student.name) or student.name
        student.course = request.form.get('course', student.course) or student.course
        if 'studying_year' in request.form:
            student.studying_year = (request.form.get('studying_year') or '').strip()
        if 'branch' in request.form:
            student.branch = (request.form.get('branch') or '').strip()
        student.registered_new_course = request.form.get('registered_new_course', student.registered_new_course or '')
        student.online_course = request.form.get('online_course', student.online_course or '')
        student.event_participation = request.form.get('event_participation', student.event_participation or '')
        student.additional_description = request.form.get('description', student.additional_description or '')
        student.custom_advisory = request.form.get('custom_advisory', student.custom_advisory or '')

        # Save results from faculty edit modal
        result_subjects = request.form.getlist('result_subject[]')
        result_grades = request.form.getlist('result_grade[]')
        results = [{'subject': s.strip(), 'grade': g.strip()} for s, g in zip(result_subjects, result_grades) if s.strip()]
        student.results_data = json.dumps(results)

        student.last_updated = datetime.now(pytz.timezone('Asia/Kolkata')).replace(tzinfo=None)
        db.session.commit()
        flash(f'Details for {student.name or reg_num} updated successfully!', 'success')
    else:
        flash('Student not found.', 'danger')
        
    return redirect(url_for('faculty_dashboard'))

@app.route('/update_global_advisory', methods=['POST'])
def update_global_advisory():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    advisory_text = request.form.get('global_advisory', '').strip()
    global_adv = GlobalAdvisory.query.first()
    if not global_adv:
        global_adv = GlobalAdvisory(id=1, content=advisory_text)
        db.session.add(global_adv)
    else:
        global_adv.content = advisory_text
    
    db.session.commit()
    flash('Global Institutional Advisory updated for all students!', 'success')
    return redirect(url_for('faculty_dashboard'))

@app.route('/update_global_observation', methods=['POST'])
def update_global_observation():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    observation_text = request.form.get('global_observation', '').strip()
    global_obs = GlobalMentorObservation.query.first()
    if not global_obs:
        global_obs = GlobalMentorObservation(id=1, content=observation_text)
        db.session.add(global_obs)
    else:
        global_obs.content = observation_text
    
    db.session.commit()
    flash('Global Mentor Observation updated for all students!', 'success')
    return redirect(url_for('faculty_dashboard'))

@app.route('/faculty')
def faculty_dashboard():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    students = StudentDetail.query.all()
    global_adv = GlobalAdvisory.query.first()
    global_advisory_text = global_adv.content if global_adv else ""
    global_obs = GlobalMentorObservation.query.first()
    global_observation_text = global_obs.content if global_obs else ""
    
    # Simplified stats for global update
    total_att_a = 0
    count_a = 0
    for s in students:
        att = s.get_attendance()
        if 'Slot A' in att:
            total_att_a += int(att['Slot A'] or 0)
            count_a += 1
    
    stats = {
        'total_students': len(students),
        'reports_generated': 152,
        'avg_attendance_a': int(total_att_a / count_a) if count_a > 0 else 0,
        'avg_attendance_b': 92
    }

    students_json = [
        {
            'reg_num': s.reg_num,
            'name': s.name,
            'course': s.course,
            'studying_year': s.studying_year,
            'branch': s.branch,
            'attendance_data': s.attendance_data,
            'marks_data': s.marks_data,
            'slot_info': s.slot_info,
            'registered_new_course': s.registered_new_course,
            'online_course': s.online_course,
            'event_participation': s.event_participation,
            'additional_description': s.additional_description,
            'custom_advisory': s.custom_advisory,
            'photo_path': s.photo_path,
            'last_updated': s.last_updated.strftime('%d-%b-%Y %I:%M %p') if s.last_updated else 'N/A',
        }
        for s in students
    ]

    return render_template('faculty.html', students=students, students_json=students_json, stats=stats,
                           global_advisory=global_advisory_text,
                           global_observation=global_observation_text,
                           show_grades_in_ppt=(GlobalSettings.query.first().show_grades_in_ppt if GlobalSettings.query.first() else False))

@app.route('/toggle_grade_ppt', methods=['POST'])
def toggle_grade_ppt():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    settings = GlobalSettings.query.first()
    if not settings:
        settings = GlobalSettings(id=1, show_grades_in_ppt=True)
        db.session.add(settings)
    else:
        settings.show_grades_in_ppt = not settings.show_grades_in_ppt
    db.session.commit()
    return redirect(url_for('faculty_dashboard'))

@app.route('/generate_report')
def generate_report():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    students = StudentDetail.query.all()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    for student in students:
        slots = student.get_slots()
        att_data = student.get_attendance()
        marks_data = student.get_marks()
        mentor_name = "GNANAVEL V K"
        
        # --- SLIDE 1: ACADEMIC PERFORMANCE ---
        slide_layout = prs.slide_layouts[6] # Blank
        slide1 = prs.slides.add_slide(slide_layout)
        
        # Header Box removed as per request

        # Banner Table for Name/Reg/Mentor (replacing the previous textbox for better alignment)
        banner_tbl = slide1.shapes.add_table(1, 6, Inches(0.2), Inches(1.4), Inches(13.0), Inches(0.5)).table
        banner_content = [
            "Mentee Name", student.name or 'N/A',
            "Reg.NO", student.reg_num or 'N/A',
            "Mentor name", mentor_name
        ]
        
        for i, text in enumerate(banner_content):
            cell = banner_tbl.cell(0, i)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.name = 'Times New Roman'
            p.font.color.rgb = RGBColor(0, 0, 0)
            p.alignment = PP_ALIGN.LEFT if i % 2 == 0 else PP_ALIGN.CENTER

        # Student Photo with gray border (matching image 1)
        left_img = Inches(0.5)
        top_img = Inches(2.2)
        width_img = Inches(2.5)
        height_img = Inches(3.0)
        
        # Draw a gray frame
        frame = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_img - Inches(0.1), top_img - Inches(0.1), width_img + Inches(0.2), height_img + Inches(0.2))
        frame.fill.solid()
        frame.fill.fore_color.rgb = RGBColor(230, 230, 230)
        frame.line.color.rgb = RGBColor(200, 200, 200)

        photo_added = False
        try:
            if student.photo_path:
                if student.photo_path.startswith('http'):
                    req = urllib.request.Request(student.photo_path, headers={'User-Agent': 'Mozilla/5.0'})
                    with urllib.request.urlopen(req) as response:
                        image_stream = BytesIO(response.read())
                    slide1.shapes.add_picture(image_stream, left_img, top_img, width=width_img, height=height_img)
                    photo_added = True
                elif os.path.exists(student.photo_path):
                    slide1.shapes.add_picture(student.photo_path, left_img, top_img, width=width_img, height=height_img)
                    photo_added = True
        except Exception as e:
            print(f"Error adding photo to PPT: {e}")

        if not photo_added:
            rect = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_img, top_img, width_img, height_img)
            rect.fill.solid()
            rect.fill.fore_color.rgb = RGBColor(100, 100, 100)
            text_frame = rect.text_frame
            text_frame.clear()
            p = text_frame.paragraphs[0]
            p.text = "No Photo"
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)

        # --- Dynamic Marks Table: one row per slot ---
        all_slots = slots if slots else ["Slot A"]
        n_data_rows = len(all_slots)              # 1 row per slot
        total_rows  = 1 + n_data_rows             # + 1 header row
        cols = 4
        table_width  = Inches(9.5)
        # Scale height: 0.55 in per data row, min 2.0 in
        table_height = max(Inches(2.0), Inches(0.55 + n_data_rows * 0.55))
        left_tbl = Inches(3.3)
        top_tbl  = Inches(2.2)

        table = slide1.shapes.add_table(total_rows, cols, left_tbl, top_tbl, table_width, table_height).table

        # Header Row
        h_labels = ["Exam", "Total Marks", "Marks Obtained", "Class Average Mark"]
        for i, h in enumerate(h_labels):
            cell = table.cell(0, i)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(112, 173, 71)   # SIMATS Green
            p = cell.text_frame.paragraphs[0]
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.bold = True
            p.font.size = Pt(18)
            p.alignment = PP_ALIGN.CENTER

        # Data rows – one row per slot (Test 1 only)
        for row_idx, slot in enumerate(all_slots, start=1):
            s_marks      = marks_data.get(slot, {})
            test1_val    = str(s_marks.get('test1', '0') or '0')
            avg_val      = str(s_marks.get('avg',   '0') or '0')
            total_m_val  = str(s_marks.get('total_marks', '-') or '-')

            r_vals = [f"Test 1 ({slot})", total_m_val, test1_val, avg_val]
            for c_idx, val in enumerate(r_vals):
                cell = table.cell(row_idx, c_idx)
                cell.text = val
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(226, 239, 218)   # Light green
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(18)
                p.font.bold = True
                p.alignment = PP_ALIGN.CENTER

        # Footer Date
        # Removed as per request

        # --- SLIDE 2: MENTOR NOTES & ATTENDANCE ---
        slide2 = prs.slides.add_slide(slide_layout)

        # Header Box removed as per request

        # Read grade toggle setting
        ppt_settings = GlobalSettings.query.first()
        show_grades = ppt_settings.show_grades_in_ppt if ppt_settings else False

        # Main Gray Content Box
        body_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.8))
        body_box.fill.solid()
        body_box.fill.fore_color.rgb = RGBColor(245, 245, 245)
        body_box.line.color.rgb = RGBColor(200, 200, 200)

        tf_body = slide2.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.6)).text_frame
        tf_body.word_wrap = True

        # "Welcome to SIMATS ENGINEERING" with green highlight
        p = tf_body.paragraphs[0]
        p.text = "Welcome to SIMATS ENGINEERING"
        p.font.bold = True
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(0, 0, 0)
        if len(p.runs) > 0: add_highlight(p.runs[0], '00FF00')

        p = tf_body.add_paragraph()
        p.text = "Dear Parent,"
        p.font.size = Pt(18)
        p.font.bold = True
        p.space_after = Pt(10)

        p = tf_body.add_paragraph()

        low_attendance = False
        for slot in slots:
            try:
                if int(att_data.get(slot, 0)) < 80:
                    low_attendance = True
                    break
            except (ValueError, TypeError):
                pass

        if low_attendance:
            p.text = f"{student.name or 'The student'} has attendance below 80%. Please maintain the attendance % above 80%."
            p.font.color.rgb = RGBColor(0, 0, 0)
            if len(p.runs) > 0: add_highlight(p.runs[0], 'FF0000')
        else:
            p.text = f"So far {student.name or 'the student'} has maintained consistent attendance in the course."
            p.font.color.rgb = RGBColor(0, 0, 0)
            if len(p.runs) > 0: add_highlight(p.runs[0], '00FF00')

        p.font.size = Pt(18)
        p.font.bold = True

        for slot in slots:
            p = tf_body.add_paragraph()
            course_name = marks_data.get(slot, {}).get('course', '')
            if course_name:
                p.text = f"Attendance for {slot}: {course_name}: {att_data.get(slot, 0)}%"
            else:
                p.text = f"Attendance for {slot}: {att_data.get(slot, 0)}%"
            p.font.size = Pt(18)
            p.font.bold = True
            if len(p.runs) > 0: add_highlight(p.runs[0], 'FFFF00')

        # Use student-specific observation if set, otherwise fall back to global observation
        global_obs_record = GlobalMentorObservation.query.first()
        global_obs_text = global_obs_record.content if global_obs_record else 'I personally advised the student to concentrate more on study and skill development.'
        observation_to_use = student.additional_description.strip() if (student.additional_description and student.additional_description.strip()) else global_obs_text
        p = tf_body.add_paragraph()
        p.space_before = Pt(15)
        p.text = observation_to_use
        p.font.size = Pt(18)
        p.font.bold = True
        if len(p.runs) > 0: add_highlight(p.runs[0], 'FFFF00')

        if show_grades:
            # --- GRADES MODE: show subject results, skip new course & event ---
            try:
                results = json.loads(student.results_data) if student.results_data else []
            except:
                results = []
            if results:
                p = tf_body.add_paragraph()
                p.text = "Subject Results:"
                p.font.size = Pt(18)
                p.font.bold = True
                p.space_before = Pt(10)
                if len(p.runs) > 0: add_highlight(p.runs[0], 'FFFF00')
                for r in results:
                    p = tf_body.add_paragraph()
                    p.text = f"  {r.get('subject', '')}  —  Grade: {r.get('grade', '')}"
                    p.font.size = Pt(18)
                    p.font.bold = True
                    if len(p.runs) > 0: add_highlight(p.runs[0], 'FFFF00')
            else:
                p = tf_body.add_paragraph()
                p.text = "No subject results added yet."
                p.font.size = Pt(18)
                p.font.bold = True
                p.space_before = Pt(10)
        else:
            # --- STANDARD MODE: show new course & event, skip grades ---
            p = tf_body.add_paragraph()
            p.text = f"New course: {student.registered_new_course or 'N/A'}"
            p.font.size = Pt(18)
            p.space_before = Pt(10)
            if len(p.runs) > 0: add_highlight(p.runs[0], 'FFFF00')

            p = tf_body.add_paragraph()
            p.space_before = Pt(15)
            if student.event_participation and student.event_participation.strip():
                p.text = f"Your ward participated in: {student.event_participation.strip()} and gave his very best throughout the journey. His dedication, hard work, and sincere efforts are truly appreciable."
            else:
                p.text = "We encourage your ward to actively participate in upcoming events and extracurricular activities to build their skills and gain valuable experience."
            p.font.size = Pt(18)
            p.font.bold = True
            if len(p.runs) > 0: add_highlight(p.runs[0], 'FFFF00')

        # Add green institutional advisory lines (Global or Custom Override)
        advisory_to_use = student.custom_advisory.strip() if (student.custom_advisory and student.custom_advisory.strip()) else (GlobalAdvisory.query.first().content if GlobalAdvisory.query.first() else "All students are advised to pay their 2nd-year tuition fees on time through the Viana Portal.\nAdditionally, kindly upload your recent passport-size photograph to your Viana profile at the earliest...")

        for adv_line in advisory_to_use.split('\n'):
            if adv_line.strip():
                p = tf_body.add_paragraph()
                p.text = adv_line.strip()
                p.font.size = Pt(18)
                p.font.bold = True
                if len(p.runs) > 0: add_highlight(p.runs[0], '00FF00')

        # Apply Times New Roman font to all paragraphs
        for paragraph in tf_body.paragraphs:
            paragraph.font.name = 'Times New Roman'
            if len(paragraph.runs) > 0:
                paragraph.runs[0].font.name = 'Times New Roman'

        # Footer Date
        # Removed as per request

    report_path = 'Mentor_Dashboard_Report.pptx'
    prs.save(report_path)
    return send_file(report_path, as_attachment=True)

@app.route('/clear_student/<reg_num>', methods=['POST'])
def clear_student(reg_num):
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    student = StudentDetail.query.get(reg_num)
    if student:
        student.slot_info = '[]'
        student.attendance_data = '{}'
        student.marks_data = '{}'
        student.registered_new_course = None
        student.online_course = None
        student.event_participation = None
        student.additional_description = None
        delete_photo(student.photo_path)
        student.photo_path = None
        student.last_updated = datetime.now(pytz.timezone('Asia/Kolkata')).replace(tzinfo=None)
        db.session.commit()
        flash(f'Report data for {student.name or reg_num} has been cleared.', 'success')
    else:
        flash('Student not found.', 'danger')
    
    return redirect(url_for('faculty_dashboard'))

@app.route('/remove_student/<reg_num>', methods=['POST'])
def remove_student(reg_num):
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    student = StudentDetail.query.get(reg_num)
    if student:
        delete_photo(student.photo_path)
        db.session.delete(student)
    
    # Delete the student's login account
    user = User.query.get(reg_num)
    if user:
        db.session.delete(user)
    
    db.session.commit()
    flash(f'Student {reg_num} has been permanently removed.', 'success')
    return redirect(url_for('faculty_dashboard'))

@app.route('/delete_all_reports', methods=['POST'])
def delete_all_reports():
    if 'user' not in session or session['role'] != 'faculty':
        return redirect(url_for('login'))
    
    # Only clear PPT/report fields — keep student accounts and rows intact
    students = StudentDetail.query.all()
    for student in students:
        delete_photo(student.photo_path)
        student.slot_info = '[]'
        student.attendance_data = '{}'
        student.marks_data = '{}'
        student.registered_new_course = None
        student.online_course = None
        student.event_participation = None
        student.additional_description = None
        student.photo_path = None

    db.session.commit()
    flash('All student report data has been cleared. Accounts are still active.', 'success')
    return redirect(url_for('faculty_dashboard'))


# ===========================================================================
# AI MENTOR COMMENT REPORT
# ===========================================================================

def _get_grades(student):
    """Return a list like [{'subject': 'Math', 'grade': 'A'}, ...] from student.results_data."""
    try:
        results = json.loads(student.results_data or "[]")
        return [
            {"subject": r.get("subject", "N/A"), "grade": r.get("grade", "N/A")}
            for r in results
        ]
    except Exception:
        return []

def correct_student_text(text):
    """
    Lightweight correction of common student-side typos:
    - Remove duplicate consecutive words (e.g. "Cryptography Cryptography")
    - Strip extra whitespace/symbols
    - Title-case names that are all-lowercase (heuristic)
    Returns (corrected_text, was_changed: bool)
    """
    if not text:
        return text, False

    original = text.strip()

    # Remove duplicate consecutive words (case-insensitive)
    corrected = re.sub(r'\b(\w+)\s+\1\b', r'\1', original, flags=re.IGNORECASE)

    # Collapse multiple spaces
    corrected = re.sub(r' {2,}', ' ', corrected).strip()

    # Title-case if the whole string is lowercase (names, subjects)
    if corrected == corrected.lower() and corrected:
        corrected = corrected.title()

    # Fix 5g/6g -> 5G/6G style (digit+letter combinations should be uppercase)
    def uppercase_tech(m):
        return m.group(0).upper()
    corrected = re.sub(r'\b(\d+[a-zA-Z]+|\b(?:ai|ml|iot|nlp|dsa|dbms|os|cn|se|pdsd|coa|psp|oop|oops|tcp|ip|rdbms|sql|html|css|js|php|ui|ux|api|sdk|ide|lms|erp|crm|gps|ble|rfid|otp|sms|cse|it|ece|eee|be|btech|mtech|mba|mca)\b)', uppercase_tech, corrected, flags=re.IGNORECASE)

    changed = corrected != original
    return corrected, changed




def _branch_skill_topics(branch):
    """Return a B.Tech discipline label and skill topics appropriate to a branch."""
    normalized = re.sub(r'[^a-z0-9]+', ' ', (branch or '').lower()).strip()
    words = set(normalized.split())
    branch_options = [
        (
            'computer science',
            {'computer science', 'cse', 'information technology', 'it', 'artificial intelligence',
             'ai', 'cs', 'machine learning', 'data science', 'cyber security', 'cybersecurity',
             'software', 'computer engineering'},
            ('Python', 'data structures and algorithms', 'full-stack web development', 'SQL and database design',
             'cloud and DevOps', 'cybersecurity fundamentals', 'machine learning', 'data visualization',
             'mobile app development', 'generative AI applications', 'Linux and networking', 'software testing',
             'computer vision', 'distributed systems', 'API development', 'version control and collaboration'),
        ),
        (
            'electronics and communication',
            {'electronics', 'ece', 'communication', 'vlsi', 'embedded systems', 'instrumentation'},
            ('embedded C', 'Internet of Things', 'Verilog and FPGA design', 'PCB design',
             'digital signal processing', 'RF and wireless communication', 'microcontrollers', 'robotics',
             'edge AI', 'communication networks', 'sensor interfacing', 'MATLAB and Simulink',
             'Linux device drivers', 'antenna design', 'electronics testing', 'low-power circuit design'),
        ),
        (
            'electrical engineering',
            {'electrical', 'eee', 'power engineering'},
            ('power systems analysis', 'PLC and SCADA', 'power electronics', 'renewable energy systems',
             'embedded systems', 'electrical CAD', 'smart-grid technology', 'electric vehicle systems',
             'MATLAB and Simulink', 'electrical protection and relays', 'motor drives', 'energy auditing',
             'industrial automation', 'battery management systems', 'control systems', 'high-voltage engineering'),
        ),
        (
            'mechanical engineering',
            {'mechanical', 'mechatronics', 'automobile', 'automotive', 'manufacturing'},
            ('3D CAD and parametric modelling', 'finite element analysis', 'computational fluid dynamics',
             'robotics and automation', 'additive manufacturing', 'electric vehicle systems',
             'computer-aided manufacturing', 'industrial IoT', 'materials selection', 'PLC programming',
             'thermal-system simulation', 'product design', 'quality engineering', 'sustainable manufacturing',
             'mechanism design', 'digital twins'),
        ),
        (
            'civil engineering',
            {'civil', 'construction', 'structural', 'geotechnical'},
            ('Building Information Modelling', 'AutoCAD and technical drafting', 'structural analysis',
             'Geographic Information Systems', 'land surveying and GIS', 'construction project scheduling',
             'geotechnical investigation', 'sustainable construction materials', 'hydrology modelling',
             'quantity surveying', 'reinforced-concrete design', 'transportation planning',
             'construction safety', 'environmental impact assessment', 'drone-based site mapping',
             'infrastructure asset management'),
        ),
        (
            'chemical engineering',
            {'chemical', 'polymer', 'petroleum'},
            ('process simulation', 'process safety engineering', 'MATLAB for process calculations',
             'industrial instrumentation', 'process control', 'energy-efficient plant design',
             'chemical data analysis', 'reaction engineering', 'process optimization',
             'environmental process engineering', 'plant layout design', 'materials characterization',
             'quality management', 'industrial automation', 'wastewater treatment', 'renewable fuels'),
        ),
        (
            'biotechnology',
            {'biotechnology', 'biotech', 'biomedical', 'bioinformatics', 'life science'},
            ('bioinformatics', 'Python for biological data', 'R for genomics', 'molecular diagnostics',
             'bioprocess engineering', 'laboratory quality systems', 'computational biology',
             'genome data analysis', 'biostatistics', 'tissue engineering', 'biomedical instrumentation',
             'cell-culture techniques', 'drug discovery informatics', 'protein structure analysis',
             'medical image analysis', 'biological database tools'),
        ),
    ]

    for label, aliases, topics in branch_options:
        if label == 'electronics and communication' and (
            'electrical' in words or 'eee' in words or 'power engineering' in normalized
        ):
            continue
        if any(alias in normalized if ' ' in alias else alias in words for alias in aliases):
            return label, topics

    return '', ()


def _assign_skill_suggestions(students):
    """Assign a distinct, branch-specific recommendation to each student in this report."""
    groups = {}
    for student in students:
        branch = (student.branch or '').strip()
        label, topics = _branch_skill_topics(branch)
        groups.setdefault((label, topics), []).append(student)

    suggestions = {}
    for (label, topics), group in groups.items():
        if not topics:
            continue
        pairs = [
            (topics[left], topics[right])
            for left in range(len(topics))
            for right in range(left + 1, len(topics))
        ]
        group.sort(key=lambda student: student.reg_num)
        start = sum(ord(char) for char in (group[0].reg_num if group else '')) % len(pairs)
        for index, student in enumerate(group):
            first_topic, second_topic = pairs[(start + index) % len(pairs)]
            suggestions[student.reg_num] = (
                f"For {label} B.Tech learning, consider an online course in {first_topic} and "
                f"{second_topic}, then apply them in a small hands-on project."
            )
    return suggestions


def _build_ai_payload(student, show_grades=False, faculty_context=None, skill_suggestion=''):
    """
    Compute facts in pure Python for one student.
    show_grades and faculty_context must be passed from the main thread
    (they must not be queried here) to avoid SQLAlchemy thread-safety issues.
    Returns (payload_dict, corrections_list)
    """
    att_data = student.get_attendance()
    marks_data = student.get_marks()
    corrections = []

    # --- Correct student-entered text fields ---
    def maybe_correct(field_value, field_name):
        corrected, changed = correct_student_text(field_value or '')
        if changed and field_value:
            corrections.append({
                'field': field_name,
                'original': field_value,
                'corrected': corrected
            })
        return corrected if corrected else (field_value or '')

    name = maybe_correct(student.name, 'Name')
    event = maybe_correct(student.event_participation, 'Event Participation')
    online_course = maybe_correct(student.online_course, 'Online Course')
    remark = maybe_correct(student.additional_description, 'Remark / Observation')

    # --- Compute attendance facts from valid recorded values only ---
    attendance_records = []

    for slot, att_val in att_data.items():
        try:
            att_pct = float(str(att_val).replace('%', '').strip() or 0)
        except (ValueError, TypeError):
            continue
        if not 0 <= att_pct <= 100:
            continue

        subject_name = marks_data.get(slot, {}).get('course', '') or slot
        attendance_records.append({'subject': subject_name, 'percent': att_pct})

    low_subjects = [record for record in attendance_records if record['percent'] < 80]

    faculty_context = faculty_context or {}
    global_advisory = (faculty_context.get('global_advisory') or '').strip()
    individual_advisory = (student.custom_advisory or '').strip()

    payload = {
        'reg_no': student.reg_num,
        'name': name,
        'gender': '',
        'attendance_records': low_subjects,
        'low_subjects': low_subjects,
        'event': event,
        'online_course': online_course,
        'branch': (student.branch or '').strip(),
        'skill_suggestion': skill_suggestion,
        'student_remark': remark,
        'faculty_advisory': individual_advisory or global_advisory,
        'faculty_observation': (faculty_context.get('global_observation') or '').strip(),
        'institution': 'SIMATS Engineering',
    }

    # Include grades only when the flag is explicitly passed as True
    # (flag must be read in the main thread before spawning worker threads)
    if show_grades:
        payload['grades'] = _get_grades(student)

    return payload, corrections


def _call_model(payload):
    """Unified LLM call – works with Gemini, OpenAI, or Groq.
    Builds the prompt (including grades & remark) and returns the comment.
    """
    # ---- Build extra instructions ------------------------------------------------
    extra_parts = []
    if payload.get('grades'):
        grades_txt = ', '.join(f"{g['subject']} ({g['grade']})" for g in payload['grades'])
        extra_parts.append(
            f"The student has the following grades that MUST be woven naturally into the comment: {grades_txt}."
        )
    remark = (payload.get('student_remark') or '').strip()
    if remark:
        extra_parts.append(
            f"The student has shared this personal remark/reason (e.g. for low attendance): '{remark}'. "
            f"Acknowledge it briefly and sensitively in the comment — do NOT copy it word for word, "
            f"just mention that the student has explained the situation and encourage them going forward."
        )
    advisory = (payload.get('faculty_advisory') or '').strip()
    if advisory:
        extra_parts.append(
            f"Include the meaning of this advisory in the comment: '{advisory}'. "
            "Rephrase it naturally in the mentor's own voice, correcting spelling and grammar as needed. "
            "Do not omit it or label it as faculty guidance."
        )
    skill_suggestion = (payload.get('skill_suggestion') or '').strip()
    if skill_suggestion:
        extra_parts.append(
            "Include this branch-relevant B.Tech skill/course recommendation verbatim as advice; do not imply "
            f"the student is already enrolled: {skill_suggestion}"
        )
    observation = (payload.get('faculty_observation') or '').strip()
    if observation:
        extra_parts.append(
            f"Consider this mentor observation as context: '{observation}'. "
            "Include only relevant guidance naturally in the mentor's voice, without identifying its source."
        )
    extra = ("\n" + " ".join(extra_parts)) if extra_parts else ""
    prompt = f"Input: {json.dumps(payload, ensure_ascii=False)}{extra}"

    # ---- Dispatch ---------------------------------------------------------------
    if PROVIDER == "gemini":
        response = _gemini_model.generate_content(
            prompt,
            generation_config={"max_output_tokens": 1600},
            request_options={"timeout": AI_REQUEST_TIMEOUT_SECONDS, "retry": None},
        )
        raw = response.text.strip()
    elif PROVIDER in {"openai", "groq"}:
        system_msg = (
            "You are a college mentor writing an individualized parent-facing progress report for one student. "
            "Follow the supplied parent-facing reference style. Begin exactly with 'Dear Parent,'. In the next sentence, "
            "introduce the student's progress using the student's full name and registration number. Put positive "
            "academic or activity highlights before attendance; place attendance information around the middle of the "
            "paragraph. Write 300 to 800 words across readable paragraphs, aiming for about 450 words when the "
            "student's record supports it. Use the extra length for practical study guidance, branch-relevant learning "
            "suggestions, and supportive encouragement grounded in the student's actual record; do not repeat points "
            "just to fill space or invent facts.\n"
            "Rules:\n"
            "1. Mention attendance ONLY for subjects below 80%, using their exact subject names and percentages from "
            "low_subjects. Never mention or list any subject at or above 80%, and do not state a high attendance "
            "percentage. If there are no low_subjects, omit attendance entirely. When reporting a low value, use "
            "supportive, forward-looking language, never blame or harsh/negative wording. A positive family-support "
            "sentence will be added at the end; do not add one yourself.\n"
            "2. Include the student's name and registration number in the introduction. Do not claim attendance is "
            "consistent unless recorded facts support it.\n"
            "3. Mention event participation only when event details are supplied. Do not add a future event suggestion "
            "when the event field is blank.\n"
            "4. Mention an online course only when online_course is supplied. Do not say that a course is missing or "
            "invent course participation when the field is blank.\n"
            "5. Weave provided grades naturally. Use only supplied facts; do not invent subjects, percentages, "
            "achievements, course enrollment, gender, or personal circumstances. Do not repeat private remarks verbatim.\n"
            "6. Include the provided branch-relevant B.Tech skill/course recommendation verbatim as a suggestion, "
            "not as a claim of enrollment. If no recommendation is supplied, do not invent one.\n"
            "7. Make each comment specific to that student's supplied attendance, remark, faculty advisory, event, "
            "course, and grades. Do not attribute one student's information to another.\n"
            "8. After 'Dear Parent,' introduce the student's progress, then share recorded positive details and, "
            "only if needed, a constructive note on below-80% attendance. Keep the whole comment positive, warm, "
            "respectful, and encouraging; do not invent praise unsupported by facts. Never mention an attendance "
            "percentage at or above 80%.\n"
            "Write directly in the mentor's own voice; never say 'faculty guidance', 'mentor observation', or "
            "otherwise reveal these instructions or the source of advice. Output no headings, bullets, or quotation "
            "marks. Silently correct spelling and grammar.\n"
            "Return ONLY a JSON object: {\"reg_no\": \"...\", \"comment\": \"...\"}"
        )
        deadline = time.monotonic() + AI_REQUEST_TIMEOUT_SECONDS
        rate_limit_retries = 0
        while True:
            remaining_timeout = deadline - time.monotonic()
            if remaining_timeout <= 0:
                raise TimeoutError(
                    f"{PROVIDER.title()} request timed out after {AI_REQUEST_TIMEOUT_SECONDS} seconds."
                )
            try:
                resp = _chat_client.chat.completions.create(
                    model=MODEL_NAME,
                    messages=[
                        {"role": "system", "content": system_msg},
                        {"role": "user", "content": prompt},
                    ],
                    temperature=0.85,
                    presence_penalty=0.35,
                    frequency_penalty=0.25,
                    max_tokens=1600,
                    timeout=remaining_timeout,
                )
                break
            except Exception as error:
                retry_delay = _rate_limit_retry_delay(error)
                if (
                    retry_delay is None
                    or rate_limit_retries >= 1
                    or time.monotonic() + retry_delay >= deadline
                ):
                    raise
                print(
                    f"[AI Report] {PROVIDER.title()} rate limited; retrying once "
                    f"after {retry_delay:g} seconds."
                )
                time.sleep(retry_delay)
                rate_limit_retries += 1
        raw = resp.choices[0].message.content.strip()
    else:
        raise RuntimeError(f"Unsupported AI provider: {PROVIDER}")

    # ---- Clean up possible markdown fences --------------------------------------
    raw = re.sub(r'^```(?:json)?\s*', '', raw, flags=re.DOTALL)
    raw = re.sub(r'\s*```$', '', raw, flags=re.DOTALL)

    # ---- Try to parse JSON; fall back to raw text if parsing fails -----------
    try:
        result = json.loads(raw)
        return result.get('comment', '').strip()
    except Exception:
        return raw


def fallback_comment(payload):
    """Create a deterministic comment using recorded facts, without filling data gaps."""
    name = (payload.get('name') or 'The student').strip()
    reg_no = (payload.get('reg_no') or '').strip()
    first = name.split()[0]
    variant = sum(ord(char) for char in f"{reg_no}|{name}") % 2
    parts = ["Dear Parent,"]

    introduction = f"It is a pleasure to share the progress of {name}"
    if reg_no:
        introduction += f" (Reg. No. {reg_no})"
    parts.append(f"{introduction}.")

    event = (payload.get('event') or '').strip()
    if event:
        event_phrases = [
            f"The recorded event participation for {first} is {event}.",
            f"{first}'s recorded event participation includes {event}.",
        ]
        parts.append(event_phrases[variant])

    online_course = (payload.get('online_course') or '').strip()
    if online_course:
        course_phrases = [
            f"The online course recorded for {first} is {online_course}.",
            f"{first} has the following online course recorded: {online_course}.",
        ]
        parts.append(course_phrases[variant])

    grades = payload.get('grades') or []
    if grades:
        grades_text = ', '.join(f"{grade['subject']} ({grade['grade']})" for grade in grades)
        parts.append(f"Recorded grades include {grades_text}.")

    low_subjects = payload.get('low_subjects') or []
    if low_subjects:
        low_text = ', '.join(
            f"{record['subject']} ({float(record['percent']):g}%)"
            for record in low_subjects
        )
        parts.append(
            f"With regular attendance in {low_text}, {first} can make the most of upcoming classes "
            "and continue progressing."
        )

    advisory = (payload.get('faculty_advisory') or '').strip()
    if advisory:
        punctuation = '' if advisory.endswith(('.', '!', '?')) else '.'
        advisory = advisory[0].upper() + advisory[1:]
        parts.append(f"{advisory}{punctuation}")

    skill_suggestion = (payload.get('skill_suggestion') or '').strip()
    if skill_suggestion:
        parts.append(skill_suggestion)

    if len(parts) == 2:
        parts.append("No additional student-specific details were available for this comment.")

    parts.append(f"We look forward to supporting {name}'s continued academic progress.")
    parts.append(_parent_cooperation_note(payload))
    return ' '.join(part for part in parts if part).strip()


def _parent_cooperation_note(payload):
    if not payload.get('low_subjects'):
        return ''
    first = (payload.get('name') or 'The student').split()[0]
    return (
        f"The encouragement and cooperation of the parents/guardians will support {first} "
        "in maintaining regular attendance and continuing this progress."
    )


def _ensure_parent_cooperation_note(comment, payload):
    note = _parent_cooperation_note(payload)
    if note and 'encouragement and cooperation of the parents/guardians' not in comment.lower():
        return f"{comment.rstrip()} {note}"
    return comment


def _validate_ai_comment(comment, payload):
    if not comment:
        raise ValueError("The AI provider returned an empty comment.")
    word_count = len(comment.split())
    if not 300 <= word_count <= 800:
        raise ValueError(
            f"The AI comment contained {word_count} words; expected between 300 and 800."
        )
    if not comment.lstrip().lower().startswith("dear parent,"):
        raise ValueError("The AI comment did not use the required parent-facing opening.")
    name = (payload.get('name') or '').strip()
    if name and name.lower() not in comment.lower():
        raise ValueError("The AI comment did not identify the student.")
    reg_no = (payload.get('reg_no') or '').strip()
    if reg_no and reg_no.lower() not in comment.lower():
        raise ValueError("The AI comment did not include the registration number.")
    for subject in payload.get('low_subjects', []):
        name = subject.get('subject', '')
        percent = subject.get('percent', -1)
        if name and name.lower() not in comment.lower():
            raise ValueError(f"Comment missing subject: {name}")
        percent_text = f"{float(percent):g}" if percent >= 0 else ''
        if percent_text and not re.search(
            rf"(?<![\d.]){re.escape(percent_text)}\s*(?:%|percent)(?![A-Za-z])",
            comment,
            flags=re.IGNORECASE,
        ):
            raise ValueError(f"Comment missing percentage for: {name}")
    skill_suggestion = (payload.get('skill_suggestion') or '').strip()
    if skill_suggestion and skill_suggestion.lower() not in comment.lower():
        raise ValueError("Comment omitted the assigned branch-specific skill recommendation.")
    allowed_attendance_percentages = {
        float(subject.get('percent', -1))
        for subject in payload.get('low_subjects', [])
    }
    for match in re.finditer(
        r'(?<![\d.])(\d+(?:\.\d+)?)\s*(?:%|percent)(?![A-Za-z])',
        comment,
        re.IGNORECASE,
    ):
        reported_percentage = float(match.group(1))
        if not any(abs(reported_percentage - allowed) < 0.01 for allowed in allowed_attendance_percentages):
            raise ValueError("The AI comment included an unsupported percentage.")


def _is_rate_limit_error(error):
    message = str(error).lower()
    return '429' in message or 'quota' in message or 'rate limit' in message


def _rate_limit_retry_delay(error):
    """Return a short retry delay for HTTP 429 responses, or None for other errors."""
    if getattr(error, 'status_code', None) != 429:
        return None

    response = getattr(error, 'response', None)
    headers = getattr(response, 'headers', {}) or {}
    delays = []
    for header in ('retry-after', 'x-ratelimit-reset-requests', 'x-ratelimit-reset-tokens'):
        value = headers.get(header)
        if value is None:
            continue
        match = re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*(?:s|sec|seconds)?\s*', str(value), re.IGNORECASE)
        if match:
            delays.append(float(match.group(1)))

    return max(delays) if delays else 1.0


def _ai_failure_message(error):
    """Return a user-safe reason for using the template fallback."""
    provider = PROVIDER.title()
    status_code = getattr(error, 'status_code', None)
    error_name = type(error).__name__.lower()

    if _is_rate_limit_error(error) or status_code == 429:
        return f"{provider} rate limit or quota reached; template fallback used."
    if status_code in (401, 403):
        return f"{provider} rejected the API credentials (HTTP {status_code}); template fallback used."
    if status_code == 404:
        return f"{provider} could not find the configured model (HTTP 404); check the model setting."
    if status_code == 400:
        return f"{provider} rejected the request (HTTP 400); check model/request settings and server logs."
    if isinstance(status_code, int) and status_code >= 500:
        return f"{provider} returned a service error (HTTP {status_code}); template fallback used."
    if 'timeout' in error_name or isinstance(error, TimeoutError):
        return f"{provider} request timed out after {AI_REQUEST_TIMEOUT_SECONDS} seconds; template fallback used."
    if isinstance(error, ValueError):
        return f"AI response did not meet report requirements ({error}); template fallback used."
    if 'connection' in error_name:
        return f"Could not connect to {provider}; check network access and server logs."
    return f"{provider} request failed ({type(error).__name__}); template fallback used. Check server logs."


def _fallback_result_for_student(student, show_grades, faculty_context, skill_suggestion, error):
    """Create one template result even if AI generation or payload preparation fails."""
    reg_no = getattr(student, 'reg_num', '') or ''
    try:
        payload, corrections = _build_ai_payload(
            student,
            show_grades=show_grades,
            faculty_context=faculty_context,
            skill_suggestion=skill_suggestion,
        )
    except Exception as payload_error:
        print(f"[AI Report] Could not prepare fallback data for {reg_no}: {payload_error}")
        payload = {
            'reg_no': reg_no,
            'name': str(getattr(student, 'name', None) or 'The student'),
            'gender': 'male',
            'low_subjects': [],
            'event': '',
            'online_course': '',
            'skill_suggestion': skill_suggestion,
        }
        corrections = []

    try:
        comment = _ensure_parent_cooperation_note(fallback_comment(payload), payload)
    except Exception as fallback_error:
        print(f"[AI Report] Template fallback failed for {reg_no}: {fallback_error}")
        name = str(payload.get('name') or 'The student')
        reg_no = payload.get('reg_no') or ''
        introduction = f"It is a pleasure to share the progress of {name}"
        if reg_no:
            introduction += f" (Reg. No. {reg_no})"
        comment = f"Dear Parent, {introduction}."
        if payload.get('low_subjects'):
            comment += f" {_parent_cooperation_note(payload)}"
        comment = comment.strip()

    return {
        'reg_no': reg_no,
        'name': payload['name'],
        'comment': comment,
        'ai_generated': False,
        'ai_warning': error,
        'corrections': corrections,
        'gender': payload['gender'],
        'low_count': len(payload['low_subjects']),
        'has_event': bool(payload.get('event', '').strip()),
    }


def _generate_for_one(student, show_grades=False, faculty_context=None, skill_suggestion=''):
    """Generate comment for a single student. Returns result dict."""
    payload, corrections = _build_ai_payload(
        student,
        show_grades=show_grades,
        faculty_context=faculty_context,
        skill_suggestion=skill_suggestion,
    )

    comment = None
    ai_generated = False
    ai_warning = None
    if AI_PROVIDER_AVAILABLE:
        try:
            comment = _call_model(payload)
            comment = _ensure_parent_cooperation_note(comment, payload)
            _validate_ai_comment(comment, payload)
            ai_generated = bool(comment)
        except Exception as error:
            print(f"[AI Report] Generation failed for {student.reg_num}: {error}")
            ai_warning = _ai_failure_message(error)
    else:
        ai_warning = AI_PROVIDER_ERROR or (
            f"{PROVIDER.title()} is unavailable; a template comment was used."
        )

    if not comment:
        if not ai_warning:
            ai_warning = "AI returned no comment; a template comment was used. Check the server logs and API configuration."
        comment = fallback_comment(payload)
        ai_generated = False
    comment = _ensure_parent_cooperation_note(comment, payload)

    return {
        'reg_no': student.reg_num,
        'name': payload['name'],
        'comment': comment,
        'ai_generated': ai_generated,
        'ai_warning': ai_warning,
        'corrections': corrections,
        'gender': payload['gender'],
        'low_count': len(payload['low_subjects']),
        'has_event': bool(payload.get('event', '').strip()),
    }


@app.route('/generate_comment_report', methods=['POST'])
def generate_comment_report():
    """
    POST /generate_comment_report
    Accepts optional JSON body { "reg_nos": ["24REG01"] } to regenerate a subset.
    Returns JSON list of { reg_no, name, comment, corrections }.
    """
    if 'user' not in session or session['role'] != 'faculty':
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json(silent=True) or {}
    reg_nos = data.get('reg_nos')  # None means all students

    if reg_nos:
        students = [StudentDetail.query.get(r) for r in reg_nos if StudentDetail.query.get(r)]
    else:
        students = StudentDetail.query.all()

    if not students:
        return jsonify({'error': 'No students found'}), 404

    # Read the grades toggle ONCE in the main thread (SQLAlchemy is not thread-safe)
    try:
        settings = GlobalSettings.query.first()
        show_grades = bool(settings and settings.show_grades_in_ppt)
    except Exception:
        show_grades = False

    global_advisory = GlobalAdvisory.query.first()
    global_observation = GlobalMentorObservation.query.first()
    faculty_context = {
        'global_advisory': global_advisory.content if global_advisory else '',
        'global_observation': global_observation.content if global_observation else '',
    }
    skill_suggestions = _assign_skill_suggestions(students)

    results = []
    # Process each UI batch concurrently so provider timeouts fit within one batch wait.
    BATCH = 3
    max_workers = min(BATCH, len(students))
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {
            executor.submit(
                _generate_for_one,
                s,
                show_grades,
                faculty_context,
                skill_suggestions.get(s.reg_num, ''),
            ): s
            for s in students
        }
        for future in as_completed(futures):
            try:
                results.append(future.result())
            except Exception as e:
                s = futures[future]
                print(f"[AI Report] Exception for {s.reg_num}: {e}")
                results.append(_fallback_result_for_student(
                    s,
                    show_grades,
                    faculty_context,
                    skill_suggestions.get(s.reg_num, ''),
                    "AI generation failed; a template comment was used. Check the server logs.",
                ))

    # Sort by name for consistent ordering
    results.sort(key=lambda x: (x['name'] or '').lower())
    return jsonify(results)


@app.route('/download_comment_report', methods=['POST'])
def download_comment_report():
    """
    POST /download_comment_report
    Accepts JSON body { "mentor_name": "GNANAVEL V K", "comments": [{ "reg_no", "name", "comment" }] }
    Returns a .docx file download.
    """
    if 'user' not in session or session['role'] != 'faculty':
        return jsonify({'error': 'Unauthorized'}), 401

    if not DOCX_AVAILABLE:
        return jsonify({'error': 'python-docx not installed on server'}), 500

    data = request.get_json(silent=True) or {}
    mentor_name = data.get('mentor_name', 'GNANAVEL V K').strip()
    comments = data.get('comments', [])

    doc = Document()

    # --- Page Setup: A4 ---
    section = doc.sections[0]
    section.page_width  = DocxInches(8.27)   # A4 width
    section.page_height = DocxInches(11.69)  # A4 height
    section.left_margin   = DocxInches(1.0)
    section.right_margin  = DocxInches(1.0)
    section.top_margin    = DocxInches(1.0)
    section.bottom_margin = DocxInches(1.0)

    # --- Helper: set run font ---
    def set_run_font(run, size_pt=11, bold=False, italic=False, colour=None):
        run.font.name = 'Arial'
        run.font.size = DocxPt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        if colour:
            run.font.color.rgb = DocxRGB(*colour)

    # --- Title ---
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run('Mentee Performance Review')
    set_run_font(title_run, size_pt=18, bold=True)

    # --- Subtitle ---
    total = len(comments)
    today_str = datetime.now().strftime('%d %B %Y')
    subtitle_para = doc.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_para.add_run(
        f"Mentor: {mentor_name}  |  Institution: SIMATS Engineering  |  Mentees: {total}  |  Date: {today_str}"
    )
    set_run_font(subtitle_run, size_pt=10, italic=True, colour=(80, 80, 80))

    doc.add_paragraph()  # spacer

    # --- Summary sentence ---
    low_att_count = sum(1 for c in comments if c.get('low_count', 0) > 0)
    event_count   = sum(1 for c in comments if c.get('has_event', False))
    summary_text = (
        f"Of the {total} mentee(s) in this group, {low_att_count} have at least one subject with attendance "
        f"below 80%, and {event_count} have notable event participation or achievements."
    )
    summary_para = doc.add_paragraph()
    summary_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    summary_run = summary_para.add_run(summary_text)
    set_run_font(summary_run, size_pt=11, italic=True)

    doc.add_paragraph()  # spacer

    # --- Per-student entries ---
    for idx, c in enumerate(comments, start=1):
        name    = (c.get('name') or 'N/A').strip()
        reg_no  = (c.get('reg_no') or '').strip()
        comment = (c.get('comment') or '').strip()

        # Heading: "1. Sayyad Ibrahim (Reg. No: 24REG01)"
        heading_para = doc.add_paragraph()
        heading_run = heading_para.add_run(f"{idx}. {name} (Reg. No: {reg_no})")
        set_run_font(heading_run, size_pt=12, bold=True)

        # Comment paragraph (justified)
        body_para = doc.add_paragraph()
        body_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body_run = body_para.add_run(comment)
        set_run_font(body_run, size_pt=11)

        doc.add_paragraph()  # spacer between students

    # --- Footer with page numbers ---
    footer = section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run()
    footer_run.font.name = 'Arial'
    footer_run.font.size = DocxPt(9)

    # Insert PAGE field into footer using XML
    fldChar_begin = DocxOxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    instrText = DocxOxmlElement('w:instrText')
    instrText.text = ' PAGE '
    fldChar_end = DocxOxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')

    footer_run._r.append(fldChar_begin)
    footer_run._r.append(instrText)
    footer_run._r.append(fldChar_end)

    # Save to in-memory buffer
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        as_attachment=True,
        download_name='Mentor_Comment_Report.docx',
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )


if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        seed_db()
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
